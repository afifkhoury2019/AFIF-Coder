"""Offline tests for cloud backends, attachments and key storage, using fake servers that
speak the Anthropic and OpenAI streaming formats and check what the app sends back.

Run: python tests/test_backends.py   (exit code 0 = all passed)
"""
import json
import os
from datetime import datetime
import sys
import tempfile
import threading
import time
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from PySide6.QtWidgets import QApplication  # noqa: E402

from afif_coder import backends  # noqa: E402
from afif_coder.agent import Agent  # noqa: E402
from afif_coder.attachments import build_message, from_image_bytes, load  # noqa: E402

app = QApplication.instance() or QApplication([])
FAILS = []
SEEN = {}          # last request body per fake server


def check(cond, what):
    print(("  ok   " if cond else "  FAIL ") + what, flush=True)
    if not cond:
        FAILS.append(what)


def sse(h, obj, event=None):
    if event:
        h.wfile.write(f"event: {event}\n".encode())
    h.wfile.write(b"data: " + json.dumps(obj).encode() + b"\n\n")
    h.wfile.flush()


# ---------------------------------------------------------------- fake Anthropic
class Anthropic(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        ok = self.headers.get("x-api-key") == "sk-ant-test"
        self.send_response(200 if ok else 401)
        self.end_headers()
        self.wfile.write(json.dumps({"data": [{"id": "claude-sonnet-5-5"}, {"id": "claude-opus-5-5"}]}
                                    if ok else {"error": {"message": "invalid x-api-key"}}).encode())

    def do_POST(self):
        b = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        SEEN["anthropic"] = b
        assert self.headers.get("x-api-key") == "sk-ant-test" and self.headers.get("anthropic-version")
        msgs = b["messages"]
        roles = [m["role"] for m in msgs]
        assert all(roles[i] != roles[i + 1] for i in range(len(roles) - 1)), roles
        n_results = sum(1 for m in msgs for c in m["content"] if isinstance(c, dict) and c.get("type") == "tool_result")
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream")
        self.end_headers()
        sse(self, {"type": "message_start", "message": {"usage": {"input_tokens": 1200, "output_tokens": 1}}}, "message_start")
        if n_results == 0:
            sse(self, {"type": "content_block_start", "index": 0, "content_block": {"type": "thinking", "thinking": ""}})
            sse(self, {"type": "content_block_delta", "index": 0, "delta": {"type": "thinking_delta", "thinking": "Read it first."}})
            sse(self, {"type": "content_block_delta", "index": 0, "delta": {"type": "signature_delta", "signature": "SIG123"}})
            sse(self, {"type": "content_block_stop", "index": 0})
            sse(self, {"type": "content_block_start", "index": 1, "content_block": {"type": "text", "text": ""}})
            sse(self, {"type": "content_block_delta", "index": 1, "delta": {"type": "text_delta", "text": "Looking at calc.py."}})
            sse(self, {"type": "content_block_start", "index": 2, "content_block": {"type": "tool_use", "id": "toolu_01", "name": "read_file", "input": {}}})
            for piece in ['{"pa', 'th": "ca', 'lc.py"}']:
                sse(self, {"type": "content_block_delta", "index": 2, "delta": {"type": "input_json_delta", "partial_json": piece}})
        elif n_results == 1:
            sse(self, {"type": "content_block_start", "index": 0, "content_block": {"type": "tool_use", "id": "toolu_02", "name": "edit_file", "input": {}}})
            sse(self, {"type": "content_block_delta", "index": 0, "delta": {"type": "input_json_delta",
                       "partial_json": json.dumps({"path": "calc.py", "old_text": "a - b", "new_text": "a + b"})}})
        else:
            sse(self, {"type": "content_block_start", "index": 0, "content_block": {"type": "text", "text": ""}})
            sse(self, {"type": "content_block_delta", "index": 0, "delta": {"type": "text_delta", "text": "Fixed add()."}})
        sse(self, {"type": "message_delta", "delta": {"stop_reason": "end_turn"}, "usage": {"output_tokens": 50}})
        sse(self, {"type": "message_stop"})


# ---------------------------------------------------------------- fake OpenAI-compatible (Gemini / DeepSeek)
class OpenAI(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def do_GET(self):
        self.send_response(200)
        self.end_headers()
        self.wfile.write(json.dumps({"data": [{"id": "models/gemini-3.8-flash"}, {"id": "models/gemini-embedding-001"},
                                              {"id": "deepseek-flash"}]}).encode())

    def do_POST(self):
        b = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        SEEN[self.server.name] = b
        tools_done = sum(1 for m in b["messages"] if m["role"] == "tool")
        self.send_response(200)
        self.send_header("Content-Type", "text/event-stream")
        self.end_headers()
        if tools_done == 0:
            if self.server.name == "deepseek":
                sse(self, {"choices": [{"index": 0, "delta": {"reasoning_content": "Need to read."}}]})
            call = {"index": 0, "id": "c1", "type": "function", "function": {"name": "read_file", "arguments": ""}}
            if self.server.name == "gemini":
                call["extra_content"] = {"google": {"thought_signature": "GSIG"}}
            sse(self, {"choices": [{"index": 0, "delta": {"tool_calls": [call]}}]})
            sse(self, {"choices": [{"index": 0, "delta": {"tool_calls": [{"index": 0, "function": {"arguments": '{"path":"calc.py"}'}}]}}]})
            sse(self, {"choices": [{"index": 0, "delta": {}, "finish_reason": "tool_calls"}]})
        else:
            sse(self, {"choices": [{"index": 0, "delta": {"content": "It subtracts."}}]})
            sse(self, {"choices": [{"index": 0, "delta": {}, "finish_reason": "stop"}]})
        sse(self, {"choices": [], "usage": {"prompt_tokens": 300, "completion_tokens": 20}})
        self.wfile.write(b"data: [DONE]\n\n")


def serve(handler, name):
    srv = ThreadingHTTPServer(("127.0.0.1", 0), handler)
    srv.name = name
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    return f"http://127.0.0.1:{srv.server_address[1]}"


def run_agent(backend, project, content):
    out = {"tools": [], "text": [], "usage": (0, 0)}
    agent = Agent({"temperature": 0.2, "max_steps": 8}, lambda: backend)
    agent.reset(project)
    agent.auto_approve = True
    agent.tool_finished.connect(lambda cid, ok, o: out["tools"].append(ok))
    agent.text_delta.connect(out["text"].append)
    agent.usage.connect(lambda i, o: out.__setitem__("usage", (i, o)))
    agent.finished.connect(lambda e: out.__setitem__("err", e))
    th = threading.Thread(target=agent.run, args=(content,))
    th.start()
    while th.is_alive():
        app.processEvents()
        time.sleep(0.02)
    app.processEvents()
    return out, agent


def project():
    d = Path(tempfile.mkdtemp())
    (d / "calc.py").write_text("def add(a, b):\n    return a - b\n")
    return d


print("Claude (Anthropic Messages API)")
a_url = serve(Anthropic, "anthropic")
backends.PROVIDERS["anthropic"].base_url = a_url
check(backends.list_models("anthropic", {"key": "sk-ant-test"}) == ["claude-sonnet-5-5", "claude-opus-5-5"], "lists models")
try:
    backends.list_models("anthropic", {"key": "wrong"})
    check(False, "bad key is reported")
except backends.BackendError as e:
    check("rejected the API key" in str(e), "bad key is reported clearly")
b = backends.make_backend("anthropic", "claude-sonnet-5-5", {"key": "sk-ant-test"})
b.URL = a_url + "/v1/messages"
proj = project()
out, agent = run_agent(b, str(proj), "fix add")
check(not out.get("err"), f"no error ({out.get('err')})")
check(out["tools"] == [True, True], "ran read_file then edit_file")
check("a + b" in (proj / "calc.py").read_text(), "file was fixed")
check(out["usage"] == (3600, 150), f"token usage summed over 3 turns {out['usage']}")
last = SEEN["anthropic"]
first_asst = next(m for m in last["messages"] if m["role"] == "assistant")
check(first_asst["content"][0].get("type") == "thinking" and first_asst["content"][0].get("signature") == "SIG123",
      "thinking block with signature sent back")
results = [c for m in last["messages"] for c in m["content"] if isinstance(c, dict) and c.get("type") == "tool_result"]
check([r["tool_use_id"] for r in results] == ["toolu_01", "toolu_02"], "tool results match tool ids")
check(last["tools"][0].get("input_schema") and "AFIF Coder" in last["system"][0]["text"], "tools + system prompt converted")
check(last["system"][0].get("cache_control") and last["tools"][-1].get("cache_control")
      and last["messages"][-1]["content"][-1].get("cache_control"), "prompt caching markers on system, tools, last turn")

print("Gemini (OpenAI-compatible)")
g_url = serve(OpenAI, "gemini")
models = backends.list_models("gemini", {"key": "g", "base_url": g_url})
check(models == ["gemini-3.8-flash"], f"lists chat models only, strips 'models/' ({models})")
png = None
from PySide6.QtCore import QBuffer, QByteArray, QIODevice  # noqa: E402
from PySide6.QtGui import QColor, QImage  # noqa: E402
img = QImage(40, 30, QImage.Format_RGB32)
img.fill(QColor("#2257A8"))
ba = QByteArray()
buf = QBuffer(ba)
buf.open(QIODevice.WriteOnly)
img.save(buf, "PNG")
png = bytes(ba)
content = build_message("what is in calc.py? see screenshot", [from_image_bytes(png)], str(proj))
b = backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": g_url})
out, agent = run_agent(b, str(project()), content)
check(not out.get("err") and out["tools"] == [True], f"tool call works ({out.get('err')})")
sent = SEEN["gemini"]["messages"]
asst = next(m for m in sent if m["role"] == "assistant")
check(asst["tool_calls"][0].get("extra_content", {}).get("google", {}).get("thought_signature") == "GSIG",
      "thought signature echoed back")
user = [m for m in sent if m["role"] == "user"][-1]
check(isinstance(user["content"], list) and user["content"][1]["type"] == "image_url"
      and "_name" not in user["content"][1], "image sent as image_url part, private keys stripped")
check(out["usage"] == (600, 40), "usage reported")

print("DeepSeek (OpenAI-compatible, reasoning)")
d_url = serve(OpenAI, "deepseek")
b = backends.make_backend("deepseek", "deepseek-flash", {"key": "d", "base_url": d_url})
out, agent = run_agent(b, str(project()), content)
sent = SEEN["deepseek"]["messages"]
asst = next(m for m in sent if m["role"] == "assistant")
check(asst.get("reasoning_content") == "Need to read.", "reasoning_content echoed back")
user = [m for m in sent if m["role"] == "user"][-1]
check(isinstance(user["content"], str) and "can't see images" in user["content"], "image turned into a note for a text-only model")

print("Switching provider mid-conversation")
b2 = backends.make_backend("anthropic", "claude-sonnet-5-5", {"key": "sk-ant-test"})
b2.URL = a_url + "/v1/messages"
system, conv = b2.convert(agent.messages)
check(conv[0]["role"] == "user" and any(c.get("type") == "tool_use" for c in conv[1]["content"]),
      "OpenAI-format history converts to Anthropic format")

print("Attachments")
tmp = Path(tempfile.mkdtemp())
(tmp / "notes.md").write_text("# Title\nhello")
import docx  # noqa: E402
import openpyxl  # noqa: E402
d = docx.Document()
d.add_paragraph("Substation 132/33 kV")
d.save(tmp / "spec.docx")
wb = openpyxl.Workbook()
wb.active.append(["Cable", "Ampacity"])
wb.active.append(["XLPE 500", 650])
wb.save(tmp / "cables.xlsx")
(tmp / "big.txt").write_text("x" * 70_000)
(tmp / "blob.bin").write_bytes(b"\x00\x01\x02" * 100)
img.save(str(tmp / "shot.png"))
files = [load(tmp / n) for n in ("notes.md", "spec.docx", "cables.xlsx", "big.txt", "blob.bin", "shot.png")]
check([f.kind for f in files] == ["text", "text", "text", "text", "error", "image"], f"kinds {[f.kind for f in files]}")
pj = project()
msg = build_message("review these", files, str(pj))
txt = msg[0]["text"]
check("Substation 132/33 kV" in txt and "XLPE 500,650" in txt and "# Title" in txt, "docx/xlsx/md text included")
check((pj / ".afif" / "attachments" / "big.txt").is_file() and "full copy saved as .afif/attachments/big.txt" in txt,
      "large file copied into project")
check(msg[1]["image_url"]["url"].startswith("data:image/"), "image becomes a data URL")

print("Screenshots reach vision models")
shot = Path(tempfile.mkdtemp())
img.save(str(shot / "ui.png"))
msgs = [{"role": "system", "content": "s"}, {"role": "user", "content": "check the UI"},
        {"role": "assistant", "content": "", "tool_calls": [{"id": "v1", "type": "function",
         "function": {"name": "view_image", "arguments": '{"path": "ui.png"}'}}]},
        {"role": "tool", "tool_call_id": "v1", "content": "Image ui.png attached below.",
         "_images": [msg[1]["image_url"]["url"]]}]
_, conv = backends.AnthropicBackend("k", "m").convert(msgs)
tr = conv[-1]["content"][0]
check(tr["type"] == "tool_result" and tr["content"][1]["type"] == "image", "Claude: image inside tool_result")
prep = backends.OpenAIBackend("u", "k", "m", vision=True)._prepare(msgs)
check(prep[-1]["role"] == "user" and prep[-1]["content"][1]["type"] == "image_url" and prep[-2]["role"] == "tool",
      "OpenAI-style: image follows tool results as a user turn")
prep = backends.OpenAIBackend("u", "k", "m", vision=False)._prepare(msgs)
check(prep[-1]["role"] == "tool" and "can't see images" in prep[-1]["content"], "text-only model gets a note")

print("Tools: web, commands, skills, images")
from afif_coder.tools import Ctx, Toolbox, ToolError, ToolResult  # noqa: E402
web = serve(type("W", (BaseHTTPRequestHandler,), {"log_message": lambda *a: None, "do_GET": lambda h: (
    h.send_response(200), h.send_header("Content-Type", "text/html"), h.end_headers(),
    h.wfile.write("<html><head><title>AFIF</title><script>x()</script></head><body><h1>Feeds</h1>"
                  "<p>BBC <a href='https://bbc.co.uk'>link</a></p><ul><li>one</li></ul></body></html>".encode()))}), "web")
tb = Toolbox(str(shot), {}, {})
page = tb.run("web_fetch", {"url": web + "/x"})
check("Title: AFIF" in page and "# Feeds" in page and "x()" not in page and "(https://bbc.co.uk)" in page,
      "web_fetch turns HTML into readable text")
lines = []
cmd = ('python -c "import time; [print(\'tick\', i, flush=True) or time.sleep(0.2) for i in range(3)]"'
       if sys.platform == "win32" else "for i in 0 1 2; do echo tick $i; sleep 0.2; done")
out = tb.run("run_command", {"command": cmd}, Ctx(on_output=lines.append))
check("[exit code 0]" in out and len([l for l in lines if "tick" in l]) == 3,
      "run_command streams live output" + ("" if "[exit code 0]" in out else f" -> {out[:300]!r} {lines!r}"))
slow = "ping -n 30 127.0.0.1" if sys.platform == "win32" else "sleep 30"
t0 = time.time()
out = tb.run("run_command", {"command": slow, "timeout_seconds": 5})
check("timed out" in out and time.time() - t0 < 15, "run_command timeout kills the process")
stop_at = time.time() + 2
t0 = time.time()
out = tb.run("run_command", {"command": slow}, Ctx(cancelled=lambda: time.time() > stop_at))
check("stopped by the user" in out and time.time() - t0 < 10, "Stop kills a running command")
r = tb.run("view_image", {"path": "ui.png"})
check(isinstance(r, ToolResult) and r.images[0].startswith("data:image/"), "view_image returns the picture")
try:
    tb.run("view_image", {"path": "../../etc/passwd"})
    check(False, "paths outside the project are refused")
except ToolError:
    check(True, "paths outside the project are refused")

from afif_coder import skills as skillmod  # noqa: E402
os.environ["AFIF_DATA"] = tempfile.mkdtemp()
sk_dir = skillmod.user_skills_dir() / "demo"
sk_dir.mkdir(parents=True)
(sk_dir / "SKILL.md").write_text("---\nname: demo-skill\ndescription: >\n  Builds demo apps\n  for tests\n---\n# Demo\nStep one.")
(sk_dir / "ref.md").write_text("reference text")
found = skillmod.discover()
check("demo-skill" in found and found["demo-skill"].description == "Builds demo apps for tests", "skill discovered with folded description")
tb2 = Toolbox(str(shot), found, {})
check("Step one." in tb2.run("use_skill", {"name": "demo-skill"}) and "ref.md" in tb2.run("use_skill", {"name": "demo-skill"}),
      "use_skill returns instructions and file list")
check(tb2.run("read_skill_file", {"skill": "demo-skill", "path": "ref.md"}) == "reference text", "read_skill_file")

print("MCP client")
fake_mcp = Path(tempfile.mkdtemp()) / "fake_mcp.py"
fake_mcp.write_text('''import json, sys
def send(m): sys.stdout.write(json.dumps(m) + "\\n"); sys.stdout.flush()
for line in sys.stdin:
    m = json.loads(line)
    if m.get("method") == "initialize":
        send({"jsonrpc": "2.0", "id": m["id"], "result": {"protocolVersion": "2025-06-18", "capabilities": {"tools": {}},
              "serverInfo": {"name": "fake", "version": "1"}}})
        send({"jsonrpc": "2.0", "id": 99, "method": "ping"})
    elif m.get("method") == "tools/list" and not m["params"].get("cursor"):
        send({"jsonrpc": "2.0", "id": m["id"], "result": {"tools": [{"name": "inspect_android_package",
              "description": "Inspect an APK", "inputSchema": {"type": "object", "properties": {"path": {"type": "string"}}}}],
              "nextCursor": "p2"}})
    elif m.get("method") == "tools/list":
        send({"jsonrpc": "2.0", "id": m["id"], "result": {"tools": [{"name": "patch_binary", "inputSchema": {"type": "object"}}]}})
    elif m.get("method") == "tools/call":
        a = m["params"]["arguments"]
        send({"jsonrpc": "2.0", "id": m["id"], "result": {"content": [{"type": "text", "text": "package com.afiftech.news from " + a.get("path", "")}],
              "isError": False}})
''')
from afif_coder.mcp import MCPServer  # noqa: E402
srv = MCPServer("rea", [sys.executable, str(fake_mcp)]).start()
check([t["name"] for t in srv.tools] == ["inspect_android_package", "patch_binary"], "initialize + paginated tools/list")
text, imgs, err = srv.call("inspect_android_package", {"path": "app.apk"})
check(text == "package com.afiftech.news from app.apk" and not err, "tools/call returns text")
sch = srv.schemas("rea")
check(sch[0]["function"]["name"] == "rea__inspect_android_package" and sch[1]["function"]["parameters"]["properties"] == {},
      "MCP tools exposed to models with a prefix")
from afif_coder.agent import READ_ONLY_MCP  # noqa: E402
check(READ_ONLY_MCP.match("inspect_android_package") and not READ_ONLY_MCP.match("patch_binary"),
      "read-only REA tools run freely, others need approval")
srv.stop()

print("Continuity: auto skills, lessons, progress, saved chats")
bundled = Path(__file__).resolve().parent.parent / "skills"
os.environ["AFIF_DATA"] = tempfile.mkdtemp()
import afif_coder.skills as sm  # noqa: E402
sm.bundled_skills_dir = lambda: bundled
allsk = sm.discover()
check({"afif-app-rebuild", "reverse-engineer-anything", "afif-word", "afif-ui-check"} <= set(allsk), f"bundled skills found ({len(allsk)})")
rb, rea = allsk["afif-app-rebuild"], allsk["reverse-engineer-anything"]
check(rb.matches("Clone the Nabd app with my branding") and rb.matches("أريد تطبيق مثل نبض") and not rb.matches("fix this bug"),
      "playbook triggers on clone requests (English and Arabic)")
check(rea.matches("decompile my old APK") and not rea.matches("I already created the area"), "REA triggers without false hits on 'rea'")
check(rea.mcp == "rea", "REA skill connects the REA server")
pj = project()
agent = Agent({"temperature": 0.2, "max_steps": 6}, lambda: b)
agent.reset(str(pj))
check("no progress notes yet" in agent.messages[0]["content"] and "afif-app-rebuild:" in agent.messages[0]["content"],
      "system prompt lists skills and asks for progress notes")
agent._auto_skills("please rebuild the Nabd news app as AFIF News")
check(agent.loaded == ["afif-app-rebuild"] and "Ground rules" in agent.messages[-1]["content"], "matching skill loaded into the chat")
agent._auto_skills("rebuild it again")
check(agent.loaded == ["afif-app-rebuild"], "skill isn't loaded twice")
tb3 = Toolbox(str(pj), allsk, {})
tb3.run("record_lesson", {"skill": "afif-app-rebuild", "lesson": "Ammon rss.php returns HTML; use Google News search feed instead."})
check("Ammon rss.php returns HTML" in tb3.run("use_skill", {"name": "afif-app-rebuild"}), "recorded lesson appears for every model")
tb3.run("update_progress", {"content": "# Progress\n## Status\nSpec done, repo created.\n## Next steps\nFirst CI build."})
from afif_coder.project import Project  # noqa: E402
P = Project(str(pj))
P.set_pinned(["afif-app-rebuild"])
agent.reset(str(pj))
sp = agent.messages[0]["content"]
check("Spec done, repo created." in sp and 'pinned="true"' in sp and "Ammon rss.php" in sp,
      "new chat sees progress notes, pinned skill and its lessons")
agent.messages.append({"role": "user", "content": "hello"})
agent.messages.append({"role": "assistant", "content": "hi"})
agent.loaded = ["afif-word"]
agent.save()
again = Agent({}, None)
again.reset(str(pj))
shown = again.resume(P.load_chat())
check([m["content"] for m in shown] == ["hello", "hi"] and again.loaded == ["afif-word"], "saved chat resumes with its loaded skills")
check((pj / ".afif" / ".gitignore").read_text().startswith("chats/"), "chats are kept out of git")
again.new_chat()
check(P.load_chat() is None and list((pj / ".afif" / "chats").glob("chat-*.json")), "new chat archives the old one")

print("Speed")
check(SEEN["anthropic"].get("output_config") == {"effort": "medium"}, "Claude gets effort from the Speed setting (balanced = medium)")
check(SEEN["gemini"].get("reasoning_effort") == "medium", "Gemini gets reasoning_effort")
check("reasoning_effort" not in SEEN["deepseek"], "DeepSeek isn't sent options it doesn't use")
thought = []
b = backends.make_backend("anthropic", "claude-sonnet-5-5", {"key": "sk-ant-test"})
b.URL = a_url + "/v1/messages"
ag2 = Agent({"temperature": 0.2, "max_steps": 6, "speed": "fast"}, lambda: b)
ag2.reset(str(project()))
ag2.auto_approve = True
ag2.thinking_delta.connect(thought.append)
times = []
ag2.timing.connect(lambda f, t: times.append(f))
th = threading.Thread(target=ag2.run, args=("fix add",)); th.start()
while th.is_alive():
    app.processEvents(); time.sleep(0.02)
app.processEvents()
check(SEEN["anthropic"]["output_config"] == {"effort": "low"}, "Speed setting Fast = low effort")
check("".join(thought) == "Read it first.", "Claude's thinking is shown live")
check(times and 0 < times[0] < 5, f"time to first output measured ({times})")

class Slow(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(200); self.send_header("Content-Type", "text/event-stream")
        chunked = self.path.startswith("/c")
        if chunked: self.send_header("Transfer-Encoding", "chunked")
        else: self.send_header("Connection", "close")
        self.end_headers()
        def w(b):
            if chunked: self.wfile.write(f"{len(b):x}\r\n".encode() + b + b"\r\n")
            else: self.wfile.write(b)
            self.wfile.flush()
        w(b'data: {"choices":[{"index":0,"delta":{"content":"Hello"}}]}\n\n')
        time.sleep(1.5)
        w(b'data: {"choices":[{"index":0,"delta":{"content":" there"}}]}\n\ndata: [DONE]\n\n')
        if chunked: self.wfile.write(b"0\r\n\r\n"); self.wfile.flush()
        else: self.close_connection = True
slow = serve(Slow, "slow")
for path, kind in (("/c/v1", "chunked"), ("/p/v1", "plain")):
    ob = backends.OpenAIBackend(slow + path, "", "m", vision=False)
    stamps, t0 = [], time.time()
    turn = ob.stream([{"role": "user", "content": "hi"}], [], None, lambda t: stamps.append(time.time() - t0), lambda: False)
    check(turn.msg["content"] == "Hello there" and stamps[0] < 0.8 and stamps[-1] > 1.2,
          f"text appears immediately, not at the end ({kind}: first {stamps[0]:.2f}s, last {stamps[-1]:.2f}s)")

print("Local engine speed")
class LocalFake(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        body = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        SEEN.setdefault("local", []).append(body)
        self.send_response(200); self.send_header("Content-Type", "application/json"); self.end_headers()
        self.wfile.write(b'{"choices":[{"message":{"content":"."}}]}')
loc = serve(LocalFake, "local")
lb = backends.LocalBackend(loc, 8192)
agl = Agent({"speed": "fast"}, lambda: lb)
agl.reset(str(project()))
agl.warm_up(lb)
w = SEEN["local"][-1]
check(w["max_tokens"] == 1 and w["stream"] is False and w["cache_prompt"] is True, "warm-up pre-reads the prompt (1 token)")
check(w["chat_template_kwargs"]["enable_thinking"] is False, "local thinking off unless Thorough")
check("Project folder" in w["messages"][0]["content"] and len(w["messages"][0]["content"]) < 3000,
      f"local models get the short system prompt ({len(w['messages'][0]['content'])} chars)")
names = [t["function"]["name"] for t in w["tools"]]
check("view_image" not in names and "edit_file" in names and "update_progress" in names and
      all("description" not in p for t in w["tools"] for p in t["function"]["parameters"]["properties"].values()),
      f"local models get a compact tool list ({len(names)} tools)")
lb.speed = "thorough"
check(lb.extra_body()["chat_template_kwargs"]["enable_thinking"] is True, "Thorough lets local models think")

print("Stop never freezes the window")
class Hang(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(200); self.send_header("Content-Type", "text/event-stream")
        self.send_header("Transfer-Encoding", "chunked"); self.end_headers()
        ev = b'data: {"choices":[{"index":0,"prompt_progress":null,"delta":{}}],"prompt_progress":{"total":8000,"cache":0,"processed":2000,"time_ms":5}}\n\n'
        self.wfile.write(f"{len(ev):x}\r\n".encode() + ev + b"\r\n"); self.wfile.flush()
        time.sleep(30)                                  # engine busy reading a long prompt
hang = serve(Hang, "hang")
hb = backends.LocalBackend(hang, 8192)
ah = Agent({"speed": "fast"}, lambda: hb)
ah.reset(str(project()))
prog, done = [], {}
ah.progress.connect(lambda d, t: prog.append((d, t)))
ah.finished.connect(lambda e: done.setdefault("err", e))
th = threading.Thread(target=ah.run, args=("hi",)); th.start()
t_end = time.time() + 3
while time.time() < t_end:
    app.processEvents(); time.sleep(0.05)
t0 = time.time(); ah.stop(); stop_took = time.time() - t0
t_end = time.time() + 5
while th.is_alive() and time.time() < t_end:
    app.processEvents(); time.sleep(0.05)
app.processEvents()
check(stop_took < 0.3, f"Stop returns instantly ({stop_took:.2f}s)")
check(not th.is_alive() and done.get("err") == "Stopped.", f"the reply ends right away when stopped ({done})")
check(prog and prog[0] == (2000, 8000), f"reading progress reported ({prog[:1]})")

import socket as _socket  # noqa: E402
_orig_shutdown = _socket.socket.shutdown
_socket.socket.shutdown = lambda self, how: None     # act like Windows: shutdown doesn't wake a blocked read
ah2 = Agent({"speed": "fast"}, lambda: backends.LocalBackend(hang, 8192))
ah2.reset(str(project()))
done2 = {}
ah2.finished.connect(lambda e: done2.setdefault("err", e))
th = threading.Thread(target=ah2.run, args=("hi",)); th.start()
t_end = time.time() + 2
while time.time() < t_end:
    app.processEvents(); time.sleep(0.05)
t0 = time.time(); ah2.stop()
while th.is_alive() and time.time() - t0 < 5:
    app.processEvents(); time.sleep(0.05)
app.processEvents()
_socket.socket.shutdown = _orig_shutdown
check(not th.is_alive() and done2.get("err") == "Stopped." and time.time() - t0 < 1.5,
      f"Stop ends the reply at once even when the OS doesn't wake the socket ({time.time() - t0:.2f}s, {done2})")

print("Local models don't re-read long skills")
al = Agent({}, lambda: hb)
al.reset(str(project()))
al._lean = True
al._auto_skills("can you clone apps")
check(len(al.messages[-1]["content"]) < 300 and "use_skill" in al.messages[-1]["content"], "local: skill suggested in one line, not 6,000 tokens")
al.messages.append({"role": "user", "content": '[AFIF Coder loaded the skill "afif-app-rebuild" because it matches this request. Follow it for this task.]\n' + "x" * 20000})
al._shrink_skill_notes()
check(len(al.messages[-1]["content"]) < 300, "an earlier full skill in a saved chat is shrunk for local models")
ac = Agent({}, None); ac.reset(str(project())); ac._lean = False
ac._auto_skills("can you clone apps")
check(len(ac.messages[-1]["content"]) > 5000, "cloud models still get the full skill")

print("Errors in Gemini's list format")
class GErr(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(429); self.send_header("Content-Type", "application/json"); self.end_headers()
        self.wfile.write(json.dumps([{"error": {"code": 429, "message": "Resource has been exhausted (e.g. check quota).",
                                                 "status": "RESOURCE_EXHAUSTED"}}]).encode())
ge = serve(GErr, "gerr")
gb = backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": ge})
out, _ = run_agent(gb, str(project()), "hi")
check("quota" in (out.get("err") or "") and "AttributeError" not in (out.get("err") or ""),
      f"a Gemini error shows its real message ({out.get('err')})")

print("Debug report")
from afif_coder import diag, report  # noqa: E402
os.environ["AFIF_DATA"] = tempfile.mkdtemp()
diag.event("probe", key="sk-ant-api03-SECRETSECRETSECRET", note="Bearer abcdefghijklmnopqrstuvwx", tokens_in=5,
           body='{"access_token": "ya29.abcdefghijklmnopqrstuvwxyz"}', gkey="AIzaSyA1234567890abcdefghijklmnop")
raw = (diag.logs_dir() / f"afif-{datetime.now():%Y-%m-%d}.jsonl").read_text()
check("SECRET" not in raw and "abcdefghijklmnopqrstuvwx" not in raw and "ya29" not in raw and "AIzaSy" not in raw
      and '"tokens_in": 5' in raw, "API keys and tokens never reach the log")
diag.install_crash_hooks()
t = threading.Thread(target=lambda: 1 / 0, name="worker-x"); t.start(); t.join()
ev = diag.read_events()
check(any(e["kind"] == "crash" and "ZeroDivisionError" in e.get("error", "") and "worker-x" in e.get("where", "")
          for e in ev), "crashes in background threads are recorded with a traceback")
b = backends.make_backend("anthropic", "claude-sonnet-5-5", {"key": "sk-ant-test"})
b.URL = a_url + "/v1/messages"
run_agent(b, str(project()), "fix add")
gb = backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": ge})
run_agent(gb, str(project()), "hi")
ev = diag.read_events()
msgs = [e for e in ev if e["kind"] == "message"]
turns = [e for e in ev if e["kind"] == "turn"]
tools = [e for e in ev if e["kind"] == "tool"]
check(len(msgs) == 2 and msgs[0]["result"] == "done" and msgs[0]["turns"] == 3 and msgs[0]["tools"] == 2
      and msgs[0]["first_words_s"] > 0, f"each message logged with timing ({msgs[0] if msgs else None})")
check(turns and "first_event_s" in turns[0]["request"] and turns[0]["request"]["status"] == 200, "each model reply logged with latency")
check([t["tool"] for t in tools[:2]] == ["read_file", "edit_file"] and all("seconds" in t for t in tools), "each tool step logged with duration")
gerr = [e for e in ev if e["kind"] == "message_error"]
check(gerr and gerr[-1]["request"].get("status") == 429 and "RESOURCE_EXHAUSTED" in gerr[-1]["request"].get("error_body", ""),
      "provider errors keep the HTTP status and the provider's message")
page = report.render_html(ev, diag.system_info({"speed": "fast", "cloud": {"gemini": {"key": "secret-should-hide"}}}))
check("Speed by model" in page and "claude-sonnet-5-5" in page and "rate limit" in page.lower() and "secret-should-hide" not in page,
      "report shows speed by model and errors, hides keys")
z = Path(tempfile.mkdtemp()) / "r.zip"
report.export_zip(z, {"speed": "fast"})
import zipfile as _zf  # noqa: E402
names = _zf.ZipFile(z).namelist()
check({"report.html", "events.jsonl"} <= set(names), f"export makes one zip ({names})")


print("1.3: large downloads don't crash the window")
from PySide6.QtCore import QObject as _QO, Signal as _Sig  # noqa: E402
from afif_coder.models import Downloader  # noqa: E402
dl = Downloader()
got = []
dl.progress.connect(lambda mid, d, t, sp: got.append((d, t)))
threading.Thread(target=lambda: dl.progress.emit("m", 5_000_000_000, 6_000_000_000, 1e6)).start()
t0 = time.time()
while not got and time.time() - t0 < 3:
    app.processEvents(); time.sleep(0.02)
check(got == [(5_000_000_000, 6_000_000_000)], f"progress above 2 GB arrives intact ({got})")
os.environ["AFIF_DATA"] = tempfile.mkdtemp()
from afif_coder import diag as _diag  # noqa: E402
for _ in range(50):
    try:
        raise OverflowError()
    except OverflowError as e:
        _diag.error("crash", e, where="main")
crashes = [e for e in _diag.read_events() if e["kind"] == "crash"]
check(len(crashes) == 1, f"a repeating crash is logged once, then counted ({len(crashes)} entries)")

print("1.3: Stop works while a cloud model is still connecting")
class SlowHeaders(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        time.sleep(6)
        try:
            self.send_response(200); self.end_headers()
        except OSError:
            pass
sh = serve(SlowHeaders, "slowheaders")
sb = backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": sh})
res = {}
stop_at = time.time() + 0.6
t0 = time.time()
r = sb.stream([{"role": "user", "content": "hi"}], [], 0.2, lambda t: None, lambda: time.time() > stop_at)
check(r is None and time.time() - t0 < 1.5, f"Stop before the first byte returns at once ({time.time() - t0:.2f}s)")

print("1.3: busy services are retried, used-up quotas are not")
CALLS = {"n": 0}
class Busy(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        CALLS["n"] += 1
        if CALLS["n"] == 1:
            self.send_response(503); self.send_header("Content-Type", "application/json"); self.end_headers()
            self.wfile.write(json.dumps([{"error": {"code": 503, "message": "high demand",
                                                     "details": [{"retryDelay": "1s"}]}}]).encode())
            return
        self.send_response(200); self.send_header("Content-Type", "text/event-stream"); self.end_headers()
        sse(self, {"choices": [{"index": 0, "delta": {"content": "ok"}}]})
        self.wfile.write(b"data: [DONE]\n\n")
bu = serve(Busy, "busy")
bb = backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": bu})
turn = bb.stream([{"role": "user", "content": "hi"}], [], 0.2, lambda t: None, lambda: False)
check(turn and turn.msg["content"] == "ok" and CALLS["n"] == 2, f"503 'high demand' retried after its delay ({bb.timing.get('retries')})")
class Daily(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        self.rfile.read(int(self.headers["Content-Length"]))
        self.send_response(429); self.send_header("Content-Type", "application/json"); self.end_headers()
        self.wfile.write(json.dumps({"error": {"message": "Quota exceeded: requests per day"}}).encode())
dd = serve(Daily, "daily")
db = backends.make_backend("groq", "openai/gpt-oss-120b", {"key": "k", "base_url": dd})
t0 = time.time()
try:
    db.stream([{"role": "user", "content": "hi"}], [], 0.2, lambda t: None, lambda: False); e = ""
except backends.BackendError as ex:
    e = str(ex)
check("quota" in e and time.time() - t0 < 2, f"a daily quota fails at once, no pointless waiting ({time.time() - t0:.1f}s)")

print("1.3: Gemini's special models are hidden")
for mid in ("gemini-3.5-transcribe", "gemini-3.5-flash-preview-tts", "gemini-2.5-computer-use-preview"):
    check(not backends.gemini_chat_model(mid), f"hidden: {mid}")
check(backends.gemini_chat_model("gemini-3.8-flash"), "kept: gemini-3.8-flash")

print("1.3: each turn sends only what fits (tokens and time)")
ag = Agent({}, None); ag.reset(str(project()))
ag.messages.append({"role": "user", "content": "build the weather app"})
for i in range(60):
    ag.messages.append({"role": "assistant", "content": "", "tool_calls": [
        {"id": f"c{i}", "type": "function", "function": {"name": "read_file", "arguments": json.dumps({"path": f"f{i}.kt"})}}]})
    ag.messages.append({"role": "tool", "tool_call_id": f"c{i}", "content": "x" * 3000})
ag.messages.append({"role": "user", "content": "proceed the efforts"})
full_len = len(ag.messages)
before = json.dumps(ag.messages)
msgs, tools = ag._payload(24000, False)
est = sum(Agent._tokens(m) for m in msgs) + len(json.dumps(tools)) / 3.5
check(est <= 24000, f"Gemini turn fits its 24K budget (~{est:.0f} tokens; the whole chat is ~{sum(Agent._tokens(m) for m in ag.messages):.0f})")
check(json.dumps(ag.messages) == before and len(ag.messages) == full_len, "the saved chat itself is untouched")
ids = set()
ok_pairs = True
for m in msgs:
    for c in m.get("tool_calls") or []:
        ids.add(c["id"])
    if m["role"] == "tool" and m["tool_call_id"] not in ids:
        ok_pairs = False
check(ok_pairs, "every tool result still follows the call it answers")
check(msgs[-1]["content"] == "proceed the efforts", "the newest message is always sent")
small, _ = ag._payload(6000, True)
check(sum(Agent._tokens(m) for m in small) < 6000 and "left out" in json.dumps(small),
      f"a 6K free tier (Groq, GitHub) gets the short prompt and a note about left-out history ({len(small)} msgs)")
msgs2, _ = ag._payload(24000, False)
check([m.get("content") for m in msgs2[1:]] == [m.get("content") for m in msgs[1:]], "the cut point is stable between turns (prompt cache stays valid)")
from afif_coder.engine import SPEED  # noqa: E402
lb = backends.LocalBackend("http://127.0.0.1:1", 32768)
SPEED["prompt_tps"] = 31.0
check(lb.budget_tokens == 4000, f"slow PC (31 tokens/s): about 4,000 tokens of chat ({lb.budget_tokens})")
SPEED["prompt_tps"] = 900.0
check(lb.budget_tokens == int(32768 * 0.6), f"fast GPU: up to 60% of the context ({lb.budget_tokens})")
SPEED["prompt_tps"] = 0.0

print("1.3: Free Auto switches when a free model is out of quota")
class Ok(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_POST(self):
        SEEN[self.server.name] = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
        self.send_response(200); self.send_header("Content-Type", "text/event-stream"); self.end_headers()
        sse(self, {"choices": [{"index": 0, "delta": {"content": "done by groq"}}]})
        self.wfile.write(b"data: [DONE]\n\n")
okurl = serve(Ok, "groqok")
backends.COOLDOWN.clear()
cands = [("Gemini Flash", lambda: backends.make_backend("gemini", "gemini-3.8-flash", {"key": "g", "base_url": dd})),
         ("Groq", lambda: backends.make_backend("groq", "openai/gpt-oss-120b", {"key": "k", "base_url": okurl}))]
rb = backends.RouterBackend(cands)
out, ag2 = run_agent(rb, str(project()), "hi")
check(not out.get("err") and "".join(out["text"]) == "done by groq", f"answered by the next free model ({out.get('err')})")
check(backends.COOLDOWN.get("Gemini Flash", 0) > time.time() + 600, "the used-up one rests (daily quota: an hour)")
check(len(json.dumps(SEEN["groqok"])) < 40000 and SEEN["groqok"]["max_tokens"] == 2048,
      f"Groq gets a prompt sized for its free limit ({len(json.dumps(SEEN['groqok']))} chars)")
out, _ = run_agent(backends.RouterBackend(cands), str(project()), "hi again")
check("done by groq" in "".join(out["text"]), "next message goes straight to the one that works")

print("1.3: tool calls fit each service")
hist = [{"role": "user", "content": "x"},
        {"role": "assistant", "content": "", "tool_calls": [{"id": "call_7_0", "type": "function",
         "function": {"name": "read_file", "arguments": "{}"}, "extra_content": {"google": {"thought_signature": "S"}}}]},
        {"role": "tool", "tool_call_id": "call_7_0", "content": "y"}]
gq = backends.make_backend("groq", "m", {"key": "k"})._prepare(hist)
check("extra_content" not in gq[1]["tool_calls"][0], "Gemini's signature isn't sent to Groq")
ms = backends.make_backend("mistral", "m", {"key": "k"})._prepare(hist)
check(len(ms[1]["tool_calls"][0]["id"]) == 9 and ms[1]["tool_calls"][0]["id"] == ms[2]["tool_call_id"],
      "Mistral gets 9-character call ids, matched in the results")
hist[1]["tool_calls"][0].pop("extra_content")
gm = backends.make_backend("gemini", "m", {"key": "k"})._prepare(hist)
check(gm[1]["tool_calls"][0]["extra_content"]["google"]["thought_signature"], "Gemini gets a signature on calls made by other models")

print("1.3: free providers")
class Lists(BaseHTTPRequestHandler):
    def log_message(self, *a): pass
    def do_GET(self):
        self.send_response(200); self.end_headers()
        self.wfile.write(json.dumps({"data": [
            {"id": "qwen/qwen3-coder:free", "supported_parameters": ["tools", "temperature"]},
            {"id": "meta/llama:free", "supported_parameters": ["temperature"]},
            {"id": "openai/gpt-5", "supported_parameters": ["tools"]},
            {"id": "whisper-large-v3"}, {"id": "llama-3.3-70b-versatile"}]}).encode())
lu = serve(Lists, "lists")
orm = backends.list_models("openrouter", {"key": "k", "base_url": lu})
check(orm == ["qwen/qwen3-coder:free"], f"OpenRouter: only free models that can use tools ({orm})")
gqm = backends.list_models("groq", {"key": "k", "base_url": lu})
check("whisper-large-v3" not in gqm and "llama-3.3-70b-versatile" in gqm, "Groq: speech models hidden")
check(backends.PROVIDERS["github"].keyless and backends.PROVIDERS["github"].free, "GitHub Models uses the GitHub sign-in")
try:
    backends.github_token = lambda: ""
    backends.make_backend("github", "openai/gpt-4.1", {}); e = ""
except backends.BackendError as ex:
    e = str(ex)
check("Sign in to GitHub" in e, "GitHub Models without a sign-in says what to do")

print("1.3: built-in light model and Claude Code")
from afif_coder import models as mdl, paths as pth, claude_code  # noqa: E402
bdir = Path(tempfile.mkdtemp())
(bdir / "Qwen3-1.7B-Q4_K_M.gguf").write_bytes(b"GGUF")
pth_old = mdl.bundled_models_dir
mdl.bundled_models_dir = lambda: bdir
found = mdl.installed_models([])
mdl.bundled_models_dir = pth_old
check(found and found[0].name == "Qwen3-1.7B-Q4_K_M.gguf", "the built-in model is listed first")
check(mdl.display_name(found[0]) == "Qwen3 1.7B (AFIF Lite)", "and named AFIF Lite")
light = [m for m in mdl.CATALOG if m["size"] <= 2.5]
check(len(light) >= 8, f"{len(light)} light models (2.5 GB or less) in the catalog")
env = claude_code.environment("http://127.0.0.1:5555")
check(env["ANTHROPIC_BASE_URL"] == "http://127.0.0.1:5555" and env["ANTHROPIC_MODEL"] == "afif-local"
      and env["ANTHROPIC_AUTH_TOKEN"] and "ANTHROPIC_API_KEY" not in env, "Claude Code is pointed at the AFIF Engine, no Anthropic key")
check(abs(claude_code.first_reply_minutes(100) - 3.33) < 0.1, "first-reply estimate shown before starting")


print("1.3.1: Claude Code runs inside the chat (no API key)")
os.environ["AFIF_DATA"] = tempfile.mkdtemp()
from afif_coder import claude_code as ccm  # noqa: E402
fake_bin = ccm.home() / "node_modules" / ".bin"
fake_bin.mkdir(parents=True, exist_ok=True)
fake = fake_bin / ("claude.cmd" if os.name == "nt" else "claude")
fake_py = fake_bin / "fake_claude.py"
fake_py.write_text(r"""
import json, os, sys
args = sys.argv[1:]
prompt = sys.stdin.read()
log = os.path.join(os.environ["CLAUDE_CONFIG_DIR"], "calls.jsonl")
os.makedirs(os.path.dirname(log), exist_ok=True)
open(log, "a").write(json.dumps({"args": args, "prompt": prompt, "base": os.environ.get("ANTHROPIC_BASE_URL"),
                                 "key": os.environ.get("ANTHROPIC_API_KEY")}) + "\n")
if "--tools" in args and os.environ.get("FAKE_NO_TOOLS_FLAG"):
    sys.stderr.write("error: unknown option '--tools'\n"); sys.exit(1)
def out(o): print(json.dumps(o), flush=True)
out({"type": "system", "subtype": "init", "session_id": "sess-1", "model": "afif-local"})
out({"type": "stream_event", "event": {"type": "message_start"}})
out({"type": "stream_event", "event": {"type": "content_block_delta", "delta": {"type": "text_delta", "text": "Reading calc.py. "}}})
out({"type": "assistant", "message": {"content": [{"type": "text", "text": "Reading calc.py. "},
     {"type": "tool_use", "id": "toolu_1", "name": "Read", "input": {"file_path": "calc.py"}}]}})
out({"type": "user", "message": {"content": [{"type": "tool_result", "tool_use_id": "toolu_1", "content": "def add(a, b): return a - b"}]}})
out({"type": "assistant", "message": {"content": [{"type": "tool_use", "id": "toolu_2", "name": "Edit",
     "input": {"file_path": "calc.py", "old_string": "a - b", "new_string": "a + b"}}]}})
open("calc.py", "w").write("def add(a, b):\n    return a + b\n")
out({"type": "user", "message": {"content": [{"type": "tool_result", "tool_use_id": "toolu_2", "content": [{"type": "text", "text": "ok"}]}]}})
out({"type": "assistant", "message": {"content": [{"type": "text", "text": "Fixed: add now adds."}]}})
out({"type": "result", "subtype": "success", "is_error": False, "result": "Fixed: add now adds.", "session_id": "sess-1",
     "num_turns": 3, "usage": {"input_tokens": 5000, "output_tokens": 80}})
""")
if os.name == "nt":
    fake.write_text(f'@"{sys.executable}" "{fake_py}" %*\n')
else:
    fake.write_text(f'#!/bin/sh\nexec "{sys.executable}" "{fake_py}" "$@"\n')
    fake.chmod(0o755)
check(ccm.installed(), "Claude Code found in AFIF Coder's data folder")
prj = project()
h = ccm.Harness("http://127.0.0.1:7777", lite=True)
out, ag = run_agent(h, str(prj), "fix add & don't break \"quotes\"")
calls = [json.loads(l) for l in (ccm.home() / "config" / "calls.jsonl").read_text().splitlines()]
check(not out.get("err") and out["tools"] == [True, True], f"Claude Code's steps appear as tool rows ({out.get('err')}, {out['tools']})")
check("".join(out["text"]) == "Reading calc.py. Fixed: add now adds.", f"text streamed once, not twice ({''.join(out['text'])!r})")
check("a + b" in (prj / "calc.py").read_text(), "the file was changed by Claude Code")
check(calls[0]["base"] == "http://127.0.0.1:7777" and calls[0]["key"] is None, "talks to the AFIF Engine, no Anthropic key")
check(calls[0]["prompt"] == "fix add & don't break \"quotes\"", "the message goes in exactly (stdin, no shell quoting)")
a0 = calls[0]["args"]
check("--tools" in a0 and "--system-prompt" in a0 and "bypassPermissions" in a0, "Lite setup and auto-approve passed on")
check(any(m.get("tool_calls") and m["tool_calls"][0]["function"]["name"] == "read_file" for m in ag.messages),
      "the chat history keeps the steps, so another model can continue")
out2, ag = run_agent(h, str(prj), "next")
calls = [json.loads(l) for l in (ccm.home() / "config" / "calls.jsonl").read_text().splitlines()]
check("--resume" in calls[-1]["args"] and calls[-1]["args"][calls[-1]["args"].index("--resume") + 1] == "sess-1",
      "the next message continues the same Claude Code session")
os.environ["FAKE_NO_TOOLS_FLAG"] = "1"
out3, _ = run_agent(ccm.Harness("http://127.0.0.1:7777", lite=True), str(project()), "hi")
os.environ.pop("FAKE_NO_TOOLS_FLAG")
check(not out3.get("err") and out3["tools"] == [True, True], f"an older Claude Code without the Lite flags falls back to the full setup ({out3.get('err')})")
print("Key storage")
enc = backends.protect("sk-secret-123")
check(enc and "sk-secret" not in enc and backends.unprotect(enc) == "sk-secret-123",
      f"round trip ({enc.split(':')[0]})")

print()
print("ALL PASSED" if not FAILS else f"{len(FAILS)} FAILED: {FAILS}")
sys.exit(1 if FAILS else 0)
