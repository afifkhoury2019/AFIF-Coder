"""Files the user attaches to a message: read them into text (or an image) for the model."""
from __future__ import annotations

import base64
import mimetypes
import shutil
from dataclasses import dataclass
from pathlib import Path

INLINE_LIMIT = 60_000          # characters of a file put directly into the message
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp"}
MAX_IMAGE_SIDE = 1568


@dataclass
class Attachment:
    path: Path
    kind: str = "text"         # "text" | "image" | "error"
    text: str = ""
    data_url: str = ""
    note: str = ""             # shown on the chip (e.g. "PDF, 12 pages")

    @property
    def name(self) -> str:
        return self.path.name


def load(path: str | Path) -> Attachment:
    p = Path(path)
    a = Attachment(p)
    try:
        size = p.stat().st_size
        ext = p.suffix.lower()
        if ext in IMAGE_EXT:
            a.kind, a.data_url, a.note = "image", _image_data_url(p), "image"
        elif ext == ".pdf":
            a.text, a.note = _pdf(p)
        elif ext == ".docx":
            a.text, a.note = _docx(p), "Word document"
        elif ext in (".xlsx", ".xlsm"):
            a.text, a.note = _xlsx(p), "Excel workbook"
        else:
            raw = p.read_bytes()[:5_000_000]
            if b"\x00" in raw[:8000]:
                a.kind, a.note = "error", "binary file, can't read"
            else:
                a.text = raw.decode("utf-8", errors="replace")
                a.note = f"{a.text.count(chr(10)) + 1:,} lines" if size < 5_000_000 else "first 5 MB"
    except Exception as e:
        a.kind, a.note = "error", f"couldn't read: {e}"
    return a


def from_image_bytes(png: bytes, name: str = "pasted-image.png") -> Attachment:
    a = Attachment(Path(name), "image", note="pasted image")
    a.data_url = "data:image/png;base64," + base64.b64encode(png).decode()
    return a


def build_message(text: str, files: list[Attachment], project: str) -> str | list:
    """Combine the typed text and attachments into message content.

    Returns plain text, or a list of parts (text + images) when an image is attached.
    """
    blocks, images = [], []
    for a in files:
        if a.kind == "image":
            images.append(a)
        elif a.kind == "text":
            body = a.text
            saved = ""
            if len(body) > INLINE_LIMIT:
                saved = _save_copy(a, project)
                body = body[:INLINE_LIMIT] + f"\n... [file continues; full copy saved as {saved}]"
            lang = a.path.suffix.lstrip(".")
            blocks.append(f"Attached file: {a.name}" + (f" (saved as {saved})" if saved else "")
                          + f"\n```{lang}\n{body}\n```")
        else:
            blocks.append(f"Attached file {a.name} could not be read ({a.note}).")
    full = text.strip()
    if blocks:
        full = (full + "\n\n" if full else "") + "\n\n".join(blocks)
    if not images:
        return full
    parts = [{"type": "text", "text": full or "See the attached image."}]
    for a in images:
        parts.append({"type": "image_url", "image_url": {"url": a.data_url}, "_name": a.name})
    return parts


# ---------------------------------------------------------------- readers
def _save_copy(a: Attachment, project: str) -> str:
    d = Path(project) / ".afif" / "attachments"
    d.mkdir(parents=True, exist_ok=True)
    dest = d / a.name
    if a.path.is_file():
        if a.path.suffix.lower() in (".pdf", ".docx", ".xlsx", ".xlsm"):
            dest = dest.with_suffix(dest.suffix + ".txt")
            dest.write_text(a.text, encoding="utf-8")
        else:
            shutil.copyfile(a.path, dest)
    else:
        dest.write_text(a.text, encoding="utf-8")
    return dest.relative_to(Path(project)).as_posix()


def _pdf(p: Path) -> tuple[str, str]:
    from pypdf import PdfReader
    r = PdfReader(str(p))
    pages = []
    for i, page in enumerate(r.pages, 1):
        t = (page.extract_text() or "").strip()
        pages.append(f"--- page {i} ---\n{t}")
        if sum(len(x) for x in pages) > 2_000_000:
            pages.append("[stopped: very large PDF]")
            break
    text = "\n\n".join(pages)
    note = f"PDF, {len(r.pages)} pages"
    if not text.replace("--- page", "").strip(" -0123456789\n"):
        note += " (no text: scanned?)"
    return text, note


def _docx(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    out = [para.text for para in d.paragraphs]
    for ti, table in enumerate(d.tables, 1):
        out.append(f"\n[Table {ti}]")
        for row in table.rows:
            out.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(out)


def _xlsx(p: Path) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"[Sheet: {ws.title}]")
        for n, row in enumerate(ws.iter_rows(values_only=True)):
            if n >= 2000:
                out.append("... [more rows]")
                break
            if any(v is not None for v in row):
                out.append(",".join("" if v is None else str(v) for v in row))
    return "\n".join(out)


def _image_data_url(p: Path) -> str:
    from PySide6.QtCore import QBuffer, QByteArray, QIODevice, Qt
    from PySide6.QtGui import QImage
    img = QImage(str(p))
    if img.isNull():
        mime = mimetypes.guess_type(p.name)[0] or "image/png"
        return f"data:{mime};base64," + base64.b64encode(p.read_bytes()).decode()
    if max(img.width(), img.height()) > MAX_IMAGE_SIDE:
        img = img.scaled(MAX_IMAGE_SIDE, MAX_IMAGE_SIDE, Qt.KeepAspectRatio, Qt.SmoothTransformation)
    ba = QByteArray()
    buf = QBuffer(ba)
    buf.open(QIODevice.WriteOnly)
    fmt = "PNG" if img.hasAlphaChannel() else "JPG"
    img.save(buf, fmt, 88)
    mime = "image/png" if fmt == "PNG" else "image/jpeg"
    return f"data:{mime};base64," + base64.b64encode(bytes(ba)).decode()
